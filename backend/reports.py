from io import BytesIO
from pathlib import Path
from datetime import date as Date
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from logic import type_map

FONT_DIR = Path(__file__).parent / "fonts"
pdfmetrics.registerFont(TTFont("Sans", str(FONT_DIR / "FreeSans.ttf")))
pdfmetrics.registerFont(TTFont("Sans-Bold", str(FONT_DIR / "FreeSansBold.ttf")))

MONTHS = ["Январь", "Февраль", "Март", "Апрель", "Май", "Июнь", "Июль", "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь"]
WD = ["Вс", "Пн", "Вт", "Ср", "Чт", "Пт", "Сб"]
BLUE = colors.HexColor("#0A84FF")
GRAY = colors.HexColor("#6E6E73")
HAIR = colors.HexColor("#E5E5EA")
ZEBRA = colors.HexColor("#F7F7FA")


def fh(h):
    t = round((h or 0) * 60)
    return f"{t // 60}:{t % 60:02d}"


def fm(v, cur="₪"):
    return f"{cur}{(v or 0):,.2f}".replace(",", " ")


def month_title(month):
    y, m = month.split("-")
    return f"{MONTHS[int(m) - 1]} {y}"


def day_rows(stats, s):
    types = type_map(s)
    rows = []
    for d in stats["days"]:
        c = d["calc"]
        if not c:
            continue
        e = d["entry"] or {}
        segs = ", ".join(f"{x['start']}–{x.get('end') or '…'}" + (f" ({int(x['rate'])}%)" if x.get("rate") else "")
                         for x in e.get("segments") or [])
        rates = ", ".join(f"{k}%: {fh(v)}" for k, v in sorted(c["breakdown"].items(), key=lambda i: float(i[0])))
        if c.get("sick_day"):
            rates = f"день болезни №{c['sick_day']} · {int(c['pay_percent'])}%"
        elif c["kind"] == "paid":
            rates = f"оплата {int(c['pay_percent'] or 0)}%"
        if c["bonus_pct"]:
            rates += f" · бонус +{c['bonus_pct']:g}%"
        rows.append({
            "date": d["date"], "wd": WD[d["weekday"]], "type": types.get(c["day_type"], {}).get("name", c["day_type"]),
            "segments": segs, "worked": c["worked_hours"], "credited": c["credited_hours"], "rates": rates,
            "gross": c["gross"], "comment": e.get("comment") or (d["holiday"] or {}).get("name", ""),
        })
    return rows


def summary_pairs(stats):
    t, tax = stats["totals"], stats["tax"]
    hours = [("Отработано", fh(t["worked_hours"])), ("Засчитано (с бонусами и отпуском)", fh(t["credited_hours"])),
             ("Норма месяца", fh(t["expected_hours"])), ("Сверхурочные", fh(t["overtime_hours"])),
             ("Рабочих дней", str(t["worked_days"])), ("Средний рабочий день", fh(t["avg_hours"]))]
    money = [("Работа по ставкам", fm(t["work_pay"])), ("Бонусы", fm(t["bonus_pay"])),
             ("Отпуск / больничный / праздники", fm(t["leave_pay"])), ("Доплаты и проезд", fm(t["extra_pay"])),
             ("Брутто", fm(t["gross"])), ("Подоходный налог", "−" + fm(tax["income_tax"])),
             ("Битуах Леуми + здоровье", "−" + fm(tax["social"])), ("Пенсия", "−" + fm(tax["pension"])),
             ("Керен иштальмут", "−" + fm(tax["study_fund"])), ("Нетто (оценка)", fm(tax["net"]))]
    return hours, money


def build_pdf(stats, s, user, cmp=()):
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=14 * mm, rightMargin=14 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm, title=f"Отчёт {stats['month']}")
    h1 = ParagraphStyle("h1", fontName="Sans-Bold", fontSize=20, leading=24)
    sub = ParagraphStyle("sub", fontName="Sans", fontSize=9.5, textColor=GRAY, leading=13)
    h2 = ParagraphStyle("h2", fontName="Sans-Bold", fontSize=12, leading=16, spaceBefore=8, spaceAfter=4)
    cell = ParagraphStyle("cell", fontName="Sans", fontSize=8, leading=10)
    story = [
        Paragraph(f"Отчёт о рабочем времени — {month_title(stats['month'])}", h1), Spacer(1, 3),
        Paragraph(f"Сотрудник: {user.name} ({user.email}) · Ставка: {fm(s['hourly_rate'])} / час · "
                  f"Норма: {s['daily_norm_hours']:g} ч/день · Сформировано: {Date.today().strftime('%d.%m.%Y')}", sub),
        Spacer(1, 10),
    ]
    hours, money = summary_pairs(stats)

    def kv(pairs, widths, bold_keys=()):
        tb = Table([[k, v] for k, v in pairs], colWidths=[w * mm for w in widths])
        st = [("FONT", (0, 0), (-1, -1), "Sans", 9), ("TEXTCOLOR", (0, 0), (0, -1), GRAY),
              ("ALIGN", (1, 0), (1, -1), "RIGHT"), ("LINEBELOW", (0, 0), (-1, -2), 0.4, HAIR),
              ("BOTTOMPADDING", (0, 0), (-1, -1), 4), ("TOPPADDING", (0, 0), (-1, -1), 4)]
        for i, (k, _) in enumerate(pairs):
            if k in bold_keys:
                st += [("FONT", (0, i), (-1, i), "Sans-Bold", 10), ("TEXTCOLOR", (0, i), (-1, i), colors.black)]
        tb.setStyle(TableStyle(st))
        return tb

    rates = sorted(stats["totals"]["pay_by_rate"].items(), key=lambda i: float(i[0]))
    rate_tb = Table([["Ставка", "Часы", "Сумма"]] + [[f"{k}%", fh(stats["totals"]["breakdown"].get(k)), fm(v)] for k, v in rates],
                    colWidths=[20 * mm, 22 * mm, 32 * mm])
    rate_tb.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "Sans", 9), ("FONT", (0, 0), (-1, 0), "Sans-Bold", 9),
                                 ("TEXTCOLOR", (0, 0), (-1, 0), GRAY), ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
                                 ("LINEBELOW", (0, 0), (-1, -1), 0.4, HAIR), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
    top = Table([[kv(hours, (54, 26)), kv(money, (68, 34), ("Брутто", "Нетто (оценка)")), rate_tb]],
                colWidths=[86 * mm, 108 * mm, 75 * mm])
    top.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
    story += [top, Spacer(1, 8)]
    if cmp:
        fmt = lambda r, v: fh(v) if r["key"] == "hours" else fm(v)  # noqa: E731
        sign = lambda r: ("+" if r["diff"] > 0 else "−" if r["diff"] < 0 else "") + fmt(r, abs(r["diff"]))  # noqa: E731
        ctb = Table([["Сверка с тлушем", "Расчёт", "Тлуш", "Разница"]] +
                    [[r["label"], fmt(r, r["calc"]), fmt(r, r["actual"]), sign(r)] for r in cmp],
                    colWidths=[60 * mm, 32 * mm, 32 * mm, 32 * mm])
        cst = [("FONT", (0, 0), (-1, -1), "Sans", 9), ("FONT", (0, 0), (-1, 0), "Sans-Bold", 9),
               ("TEXTCOLOR", (1, 0), (-1, 0), GRAY), ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
               ("LINEBELOW", (0, 0), (-1, -1), 0.4, HAIR), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]
        for i, r in enumerate(cmp, 1):
            if abs(r["diff"]) >= 1:
                good = r["diff"] < 0 if r["key"] in ("income_tax", "social", "pension") else r["diff"] > 0
                cst.append(("TEXTCOLOR", (3, i), (3, i), colors.HexColor("#34C759" if good else "#FF3B30")))
        ctb.setStyle(TableStyle(cst))
        ctb.hAlign = "LEFT"
        story += [ctb, Spacer(1, 8)]
    story += [Paragraph("Журнал дней", h2)]

    head = ["Дата", "День", "Категория", "Время", "Часы", "Засчит.", "Ставки", "Сумма", "Комментарий"]
    data = [head]
    for r in day_rows(stats, s):
        data.append([r["date"][8:] + "." + r["date"][5:7], r["wd"], Paragraph(r["type"], cell), Paragraph(r["segments"], cell),
                     fh(r["worked"]), fh(r["credited"]), Paragraph(r["rates"], cell), fm(r["gross"]), Paragraph(r["comment"], cell)])
    t = stats["totals"]
    data.append(["Итого", "", "", "", fh(t["worked_hours"]), fh(t["credited_hours"]), "", fm(t["gross"]), ""])
    widths = [14, 10, 34, 44, 14, 15, 52, 24, 62]
    tb = Table(data, colWidths=[w * mm for w in widths], repeatRows=1)
    st = [("FONT", (0, 0), (-1, -1), "Sans", 8), ("FONT", (0, 0), (-1, 0), "Sans-Bold", 8),
          ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("BACKGROUND", (0, 0), (-1, 0), BLUE),
          ("FONT", (0, -1), (-1, -1), "Sans-Bold", 8.5), ("LINEABOVE", (0, -1), (-1, -1), 0.8, colors.black),
          ("ALIGN", (4, 0), (5, -1), "RIGHT"), ("ALIGN", (7, 0), (7, -1), "RIGHT"),
          ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("BOTTOMPADDING", (0, 0), (-1, -1), 3), ("TOPPADDING", (0, 0), (-1, -1), 3)]
    for i in range(1, len(data) - 1):
        if i % 2 == 0:
            st.append(("BACKGROUND", (0, i), (-1, i), ZEBRA))
    tb.setStyle(TableStyle(st))
    story += [tb, Spacer(1, 6),
              Paragraph("Налоги рассчитаны приблизительно по шкале Израиля (подоходный налог, нкудот зикуй, Битуах Леуми, пенсия). "
                        "Итоговая сумма в тлуш маскорет может отличаться.", sub)]
    doc.build(story)
    return buf.getvalue()


def build_xlsx(stats, s, user, cmp=()):
    wb = Workbook()
    bold, head_fill = Font(bold=True), PatternFill("solid", fgColor="0A84FF")
    ws = wb.active
    ws.title = "Сводка"
    ws.append([f"Отчёт о рабочем времени — {month_title(stats['month'])}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"{user.name} ({user.email})", f"Ставка: {s['hourly_rate']} ₪/ч"])
    ws.append([])
    hours, money = summary_pairs(stats)
    for k, v in hours + [("", "")] + money:
        ws.append([k, v])
    ws.append([])
    ws.append(["Ставка", "Часы", "Сумма, ₪"])
    for c in ws[ws.max_row]:
        c.font = bold
    for k, v in sorted(stats["totals"]["pay_by_rate"].items(), key=lambda i: float(i[0])):
        ws.append([f"{k}%", round(stats["totals"]["breakdown"].get(k, 0), 2), round(v, 2)])
    ws.column_dimensions["A"].width = 38
    if cmp:
        ws.append([])
        ws.append(["Сверка с тлушем", "Расчёт", "Тлуш", "Разница"])
        for c in ws[ws.max_row]:
            c.font = bold
        for r in cmp:
            ws.append([r["label"], r["calc"], r["actual"], r["diff"]])
        ws.column_dimensions["D"].width = 14
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 14

    wd = wb.create_sheet("Дни")
    head = ["Дата", "День", "Категория", "Время", "Часы", "Засчитано, ч", "Ставки", "Сумма, ₪", "Комментарий"]
    wd.append(head)
    for c in wd[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(vertical="center")
    for r in day_rows(stats, s):
        wd.append([r["date"], r["wd"], r["type"], r["segments"], round(r["worked"], 2), round(r["credited"], 2),
                   r["rates"], round(r["gross"], 2), r["comment"]])
    t = stats["totals"]
    wd.append(["Итого", "", "", "", round(t["worked_hours"], 2), round(t["credited_hours"], 2), "", round(t["gross"], 2), ""])
    for c in wd[wd.max_row]:
        c.font = bold
    for col, w in zip("ABCDEFGHI", [12, 6, 24, 30, 9, 13, 36, 12, 40]):
        wd.column_dimensions[col].width = w
    wd.freeze_panes = "A2"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
